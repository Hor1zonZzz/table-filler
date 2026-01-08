"""Excel 导出器。

将批处理结果导出到 Excel 文件，包含多个 Sheet。
"""

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .base import ExporterBase
from ..models.recipe import Recipe


class ExcelExporter(ExporterBase):
    """Excel 导出器。

    导出结果到 Excel 文件，包含以下 Sheet：
    - Successful: 成功提取的记录
    - Needs Review: 需要人工审核的记录
    - Failed: 处理失败的记录
    - Summary: 处理统计信息
    """

    @property
    def format_name(self) -> str:
        return "excel"

    @property
    def file_extension(self) -> str:
        return ".xlsx"

    def export(
        self,
        results: dict[str, Any],
        output_path: Path,
        recipe: Recipe,
        include_failed: bool = False,
        include_needs_review: bool = True,
        **options: Any,
    ) -> dict[str, Any]:
        """导出到 Excel 文件。

        Args:
            results: 批处理结果
            output_path: 输出文件路径
            recipe: Recipe 配方
            include_failed: 是否包含失败记录，默认 False
            include_needs_review: 是否包含需审核记录，默认 True
            **options: 其他选项

        Returns:
            导出结果信息
        """
        # 获取字段名称列表
        field_names = recipe.get_field_names()

        # 确保目录存在
        output_dir = output_path.parent
        if output_dir and not output_dir.exists():
            output_dir.mkdir(parents=True, exist_ok=True)

        # 创建工作簿
        wb = Workbook()

        # Sheet 1: Successful
        ws_success = wb.active
        ws_success.title = "Successful"
        self._write_records_sheet(
            ws_success,
            results.get("successful", []),
            field_names,
            highlight_issues=False,
        )

        sheets_created = ["Successful"]

        # Sheet 2: Needs Review
        if include_needs_review and results.get("needs_review"):
            ws_review = wb.create_sheet("Needs Review")
            self._write_records_sheet(
                ws_review,
                results["needs_review"],
                field_names,
                highlight_issues=True,
            )
            sheets_created.append("Needs Review")

        # Sheet 3: Failed
        if include_failed and results.get("failed"):
            ws_failed = wb.create_sheet("Failed")
            self._write_failed_sheet(ws_failed, results["failed"])
            sheets_created.append("Failed")

        # Sheet 4: Summary
        ws_summary = wb.create_sheet("Summary")
        self._write_summary_sheet(ws_summary, results, recipe)
        sheets_created.append("Summary")

        # 保存工作簿
        try:
            wb.save(output_path)
        except Exception as e:
            return {
                "status": "error",
                "error": f"保存 Excel 文件失败: {str(e)}",
            }

        return {
            "status": "success",
            "file_path": str(output_path.absolute()),
            "message": f"Excel 文件已导出: {output_path}",
            "statistics": {
                "successful": len(results.get("successful", [])),
                "needs_review": len(results.get("needs_review", []))
                if include_needs_review
                else 0,
                "failed": len(results.get("failed", [])) if include_failed else 0,
            },
            "sheets": sheets_created,
        }

    def _write_records_sheet(
        self,
        ws,
        records: list[dict[str, Any]],
        field_names: list[str],
        highlight_issues: bool = False,
    ) -> None:
        """写入记录到工作表。"""
        # 定义样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        issue_fill = PatternFill(
            start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
        )
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # 构建表头
        headers = ["#", "文件名", "状态", "置信度"] + field_names
        if highlight_issues:
            headers.append("问题")

        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # 写入数据行
        for row_idx, record in enumerate(records, 2):
            # 行号
            ws.cell(row=row_idx, column=1, value=row_idx - 1).border = border

            # 文件名
            doc_path = record.get("document_path", record.get("pdf_path", ""))
            ws.cell(row=row_idx, column=2, value=Path(doc_path).name).border = border

            # 状态
            status = record.get("status", "UNKNOWN")
            status_cell = ws.cell(row=row_idx, column=3, value=status)
            status_cell.border = border
            if status == "PASS":
                status_cell.fill = PatternFill(
                    start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
                )
            elif status == "NEEDS_REVIEW":
                status_cell.fill = issue_fill

            # 置信度
            confidence = record.get("confidence", 0.0)
            conf_cell = ws.cell(row=row_idx, column=4, value=f"{confidence:.0%}")
            conf_cell.border = border

            # 字段值
            final_record = record.get("final_record", record.get("fields", {}))
            for col_idx, field_name in enumerate(field_names, 5):
                field_data = final_record.get(field_name, {})

                # 处理不同的数据结构
                if isinstance(field_data, dict):
                    value = field_data.get("value", "")
                    field_confidence = field_data.get("confidence", 1.0)
                else:
                    value = field_data
                    field_confidence = 1.0

                # 格式化值
                if value == "NOT_FOUND":
                    display_value = "—"
                elif value is None:
                    display_value = ""
                else:
                    display_value = str(value)

                cell = ws.cell(row=row_idx, column=col_idx, value=display_value)
                cell.border = border

                # 高亮低置信度值
                if field_confidence < 0.8 and highlight_issues:
                    cell.fill = issue_fill

            # 问题列
            if highlight_issues:
                issues = record.get("issues", [])
                if issues:
                    issue_text = "; ".join(
                        f"[{i.get('field_name', '?')}] {i.get('message', '')}"
                        for i in issues
                    )
                else:
                    issue_text = ""
                issue_cell = ws.cell(row=row_idx, column=len(headers), value=issue_text)
                issue_cell.border = border
                issue_cell.alignment = Alignment(wrap_text=True)

        # 调整列宽
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15

        # 文件名列加宽
        ws.column_dimensions["B"].width = 30

        # 问题列加宽
        if highlight_issues:
            ws.column_dimensions[get_column_letter(len(headers))].width = 50

    def _write_failed_sheet(
        self,
        ws,
        failed_records: list[dict[str, Any]],
    ) -> None:
        """写入失败记录到工作表。"""
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="C00000", end_color="C00000", fill_type="solid"
        )
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = ["#", "文件名", "错误信息"]

        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        # 写入数据
        for row_idx, record in enumerate(failed_records, 2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1).border = border

            doc_path = record.get("document_path", record.get("pdf_path", "unknown"))
            ws.cell(row=row_idx, column=2, value=Path(doc_path).name).border = border

            error = record.get("error", "未知错误")
            error_cell = ws.cell(row=row_idx, column=3, value=error)
            error_cell.border = border
            error_cell.alignment = Alignment(wrap_text=True)

        # 调整列宽
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 60

    def _write_summary_sheet(
        self,
        ws,
        results: dict[str, Any],
        recipe: Recipe,
    ) -> None:
        """写入统计摘要到工作表。"""
        header_font = Font(bold=True)
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # 标题
        title_cell = ws.cell(row=1, column=1, value="批处理统计摘要")
        title_cell.font = Font(bold=True, size=14)

        # 统计数据
        stats = [
            ("处理时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("配方名称", recipe.name),
            ("文档类型", recipe.document_type),
            ("字段数量", len(recipe.fields)),
            ("", ""),  # 空行
            ("总处理数", results.get("total_count", 0)),
            ("成功", len(results.get("successful", []))),
            ("需审核", len(results.get("needs_review", []))),
            ("失败", len(results.get("failed", []))),
        ]

        # 计算成功率
        total = results.get("total_count", 0)
        successful = len(results.get("successful", []))
        if total > 0:
            success_rate = f"{successful / total:.1%}"
        else:
            success_rate = "N/A"
        stats.append(("成功率", success_rate))

        # 写入统计数据
        for row_idx, (label, value) in enumerate(stats, 3):
            if label:
                label_cell = ws.cell(row=row_idx, column=1, value=label)
                label_cell.font = header_font
                label_cell.border = border

                value_cell = ws.cell(row=row_idx, column=2, value=value)
                value_cell.border = border

        # 调整列宽
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 25
