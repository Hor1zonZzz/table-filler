"""Excel export tool for batch processing results."""

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_to_excel(
    output_path: str | None,
    tool_context: Any,
    include_failed: bool = False,
    include_needs_review: bool = True,
) -> dict[str, Any]:
    """
    Export batch processing results to an Excel file.

    Args:
        output_path: Path for the output Excel file (optional, auto-generated if not provided)
        tool_context: ADK tool context for state access
        include_failed: Whether to include failed records in a separate sheet
        include_needs_review: Whether to include records needing review

    Returns:
        Export result with file path and statistics

    Example:
        export_to_excel("/path/to/output.xlsx")
    """
    # Get batch results from state
    batch_results = tool_context.state.get("temp:batch_results")
    if not batch_results:
        return {
            "status": "error",
            "error": "No batch results found. Run batch_process_pdfs first.",
        }

    # Get form config for field names
    form_config = tool_context.state.get("form_config", [])
    field_names = [f["name"] for f in form_config]

    # Generate output path if not provided
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"table_filler_export_{timestamp}.xlsx"

    # Ensure directory exists
    output_dir = Path(output_path).parent
    if output_dir and not output_dir.exists():
        output_dir.mkdir(parents=True, exist_ok=True)

    # Create workbook
    wb = Workbook()

    # Sheet 1: Successful records
    ws_success = wb.active
    ws_success.title = "Successful"
    _write_records_sheet(
        ws_success,
        batch_results.get("successful", []),
        field_names,
        highlight_issues=False,
    )

    # Sheet 2: Needs Review (if requested and has data)
    if include_needs_review and batch_results.get("needs_review"):
        ws_review = wb.create_sheet("Needs Review")
        _write_records_sheet(
            ws_review,
            batch_results["needs_review"],
            field_names,
            highlight_issues=True,
        )

    # Sheet 3: Failed (if requested and has data)
    if include_failed and batch_results.get("failed"):
        ws_failed = wb.create_sheet("Failed")
        _write_failed_sheet(ws_failed, batch_results["failed"])

    # Sheet 4: Summary
    ws_summary = wb.create_sheet("Summary")
    _write_summary_sheet(ws_summary, batch_results)

    # Save workbook
    try:
        wb.save(output_path)
    except Exception as e:
        return {
            "status": "error",
            "error": f"Failed to save Excel file: {str(e)}",
        }

    return {
        "status": "success",
        "file_path": str(Path(output_path).absolute()),
        "message": f"Excel file exported to: {output_path}",
        "statistics": {
            "successful": len(batch_results.get("successful", [])),
            "needs_review": len(batch_results.get("needs_review", [])) if include_needs_review else 0,
            "failed": len(batch_results.get("failed", [])) if include_failed else 0,
        },
        "sheets": [
            "Successful",
            "Needs Review" if include_needs_review and batch_results.get("needs_review") else None,
            "Failed" if include_failed and batch_results.get("failed") else None,
            "Summary",
        ],
    }


def _write_records_sheet(
    ws,
    records: list[dict[str, Any]],
    field_names: list[str],
    highlight_issues: bool = False,
) -> None:
    """Write records to a worksheet."""
    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    issue_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Build headers
    headers = ["#", "PDF File", "Status", "Confidence"] + field_names
    if highlight_issues:
        headers.append("Issues")

    # Write header row
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Write data rows
    for row_idx, record in enumerate(records, 2):
        # Row number
        ws.cell(row=row_idx, column=1, value=row_idx - 1).border = border

        # PDF file
        pdf_path = record.get("pdf_path", "")
        ws.cell(row=row_idx, column=2, value=Path(pdf_path).name).border = border

        # Status
        status = record.get("status", "UNKNOWN")
        status_cell = ws.cell(row=row_idx, column=3, value=status)
        status_cell.border = border
        if status == "PASS":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        elif status == "NEEDS_REVIEW":
            status_cell.fill = issue_fill

        # Confidence
        confidence = record.get("confidence", 0.0)
        conf_cell = ws.cell(row=row_idx, column=4, value=f"{confidence:.0%}")
        conf_cell.border = border

        # Field values
        final_record = record.get("final_record", {})
        for col_idx, field_name in enumerate(field_names, 5):
            field_data = final_record.get(field_name, {})

            # Handle different data structures
            if isinstance(field_data, dict):
                value = field_data.get("value", "")
                field_confidence = field_data.get("confidence", 1.0)
            else:
                value = field_data
                field_confidence = 1.0

            # Format value
            if value == "NOT_FOUND":
                display_value = "—"
            elif value is None:
                display_value = ""
            else:
                display_value = str(value)

            cell = ws.cell(row=row_idx, column=col_idx, value=display_value)
            cell.border = border

            # Highlight low confidence values
            if field_confidence < 0.8 and highlight_issues:
                cell.fill = issue_fill

        # Issues column
        if highlight_issues:
            issues = record.get("issues", [])
            if issues:
                issue_text = "; ".join(
                    f"[{i.get('field_name', '?')}] {i.get('message', '')}"
                    for i in issues
                )
            else:
                issue_text = ""
            issue_cell = ws.cell(row=row_idx, column=len(headers), value=issue_text)
            issue_cell.border = border
            issue_cell.alignment = Alignment(wrap_text=True)

    # Adjust column widths
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # Make PDF column wider
    ws.column_dimensions["B"].width = 30

    # Make issues column wider if present
    if highlight_issues:
        ws.column_dimensions[get_column_letter(len(headers))].width = 50


def _write_failed_sheet(ws, failed_records: list[dict[str, Any]]) -> None:
    """Write failed records to a worksheet."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["#", "PDF File", "Error"]

    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    # Write data
    for row_idx, record in enumerate(failed_records, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1).border = border

        pdf_path = record.get("pdf_path", "unknown")
        ws.cell(row=row_idx, column=2, value=Path(pdf_path).name).border = border

        error = record.get("error", "Unknown error")
        error_cell = ws.cell(row=row_idx, column=3, value=error)
        error_cell.border = border
        error_cell.alignment = Alignment(wrap_text=True)

    # Adjust column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 60


def _write_summary_sheet(ws, batch_results: dict[str, Any]) -> None:
    """Write summary statistics to a worksheet."""
    header_font = Font(bold=True)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title
    title_cell = ws.cell(row=1, column=1, value="Batch Processing Summary")
    title_cell.font = Font(bold=True, size=14)

    # Statistics
    stats = [
        ("Process Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Processed", batch_results.get("total_count", 0)),
        ("Successful", len(batch_results.get("successful", []))),
        ("Needs Review", len(batch_results.get("needs_review", []))),
        ("Failed", len(batch_results.get("failed", []))),
    ]

    # Calculate success rate
    total = batch_results.get("total_count", 0)
    successful = len(batch_results.get("successful", []))
    if total > 0:
        success_rate = f"{successful / total:.1%}"
    else:
        success_rate = "N/A"
    stats.append(("Success Rate", success_rate))

    # Write statistics
    for row_idx, (label, value) in enumerate(stats, 3):
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = header_font
        label_cell.border = border

        value_cell = ws.cell(row=row_idx, column=2, value=value)
        value_cell.border = border

    # Adjust column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 25
